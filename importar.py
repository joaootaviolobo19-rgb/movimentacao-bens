# -*- coding: utf-8 -*-
"""Importa 'Levantamento Geral1 (1).xlsx' para dados/*.json"""
import json, os, glob
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
DADOS = os.path.join(BASE, "dados")
os.makedirs(DADOS, exist_ok=True)

# localiza o Excel na pasta automaticamente
excel = None
for padrao in ["Levantamento*.xlsx", "*.xlsx"]:
    achados = glob.glob(os.path.join(BASE, padrao))
    if achados:
        excel = achados[0]; break

if not excel:
    print("ERRO: Nenhum .xlsx encontrado na pasta", BASE)
    raise SystemExit(1)

print("Lendo:", excel)
wb = load_workbook(excel, data_only=True)

# --- BENS (aba 'controle') ---
ws = wb["controle"]

# >>> CORREÇÃO: mantém APENAS a primeira ocorrência de cada cabeçalho
headers_raw = [c.value for c in ws[1]]
vistos = set()
headers = []
for h in headers_raw:
    if h is None:
        headers.append(None)
        continue
    chave = str(h).strip().lower()
    if chave in vistos:
        headers.append(None)      # duplicado → ignora
    else:
        vistos.add(chave)
        headers.append(h)

bens = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if not row or not row[0]: continue
    r = {}
    for h, v in zip(headers, row):
        if h is None: continue
        r[str(h).strip().lower()] = "" if v is None else str(v).strip()

    # ignora linhas que só têm "Chave de auditoria" preenchida (rodapé com fórmulas)
    if not r.get("categoria") and not r.get("hostname") and not r.get("patrimônio"):
        continue

    bens.append({
        "categoria":    r.get("categoria",""),
        "modelo":       r.get("modelo",""),
        "hostname":     r.get("hostname",""),
        "patrimonio":   r.get("patrimônio", r.get("patrimonio","")),
        "departamento": r.get("departamento",""),
        "responsavel":  r.get("responsável", r.get("responsavel","")),
        "marca":        r.get("marca",""),
        "serial":       r.get("serial",""),
        "ramal":        r.get("ramal",""),
        "ip":           r.get("ip",""),
        "observacao":   r.get("observação", r.get("observacao","")),
        "cpu":          r.get("cpu",""),
        "disco":        r.get("disco",""),
        "ram":          r.get("ram",""),
        "iso":          r.get("iso",""),
        "status":       r.get("status",""),
    })

for i,b in enumerate(bens, start=1): b["id"] = i

with open(os.path.join(DADOS,"bens.json"),"w",encoding="utf-8") as f:
    json.dump(bens, f, ensure_ascii=False, indent=2)
print(f"OK - {len(bens)} bens importados")

# --- FUNCIONÁRIOS (aba 'Responsável') ---
try:
    ws = wb["Responsável"]
    funcs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]: continue
        funcs.append({
            "nome":  str(row[0]).strip(),
            "cargo": str(row[1] or "").strip() if len(row)>1 else "",
            "setor": str(row[2] or "").strip() if len(row)>2 else "",
        })
    existentes = {f["nome"].upper() for f in funcs}
    for b in bens:
        r = (b.get("responsavel") or "").strip()
        if r and r.upper() not in existentes:
            funcs.append({"nome": r.title(), "cargo":"", "setor":""})
            existentes.add(r.upper())
    with open(os.path.join(DADOS,"funcionarios.json"),"w",encoding="utf-8") as f:
        json.dump(funcs, f, ensure_ascii=False, indent=2)
    print(f"OK - {len(funcs)} funcionarios importados")
except Exception as e:
    print("Aviso - erro ao ler aba Responsavel:", e)

# --- DEPARTAMENTOS (agora funciona de verdade) ---
deps = sorted({b["departamento"] for b in bens if b["departamento"]})
with open(os.path.join(DADOS,"departamentos.json"),"w",encoding="utf-8") as f:
    json.dump(deps, f, ensure_ascii=False, indent=2)
print(f"OK - {len(deps)} departamentos importados")

# --- Diagnóstico rápido ---
sem_resp = sum(1 for b in bens if not b["responsavel"])
print(f"\nDiagnostico:")
print(f"  Bens COM responsavel: {len(bens) - sem_resp}")
print(f"  Bens SEM responsavel: {sem_resp}")
print(f"\nImportacao concluida! Agora rode: python app.py")